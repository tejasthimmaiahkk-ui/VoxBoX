"""
Builds the college-format project report (St. Francis College BCA rules).

Inputs
  * the previous long-form report (VoxBox_Project_Report.previous.docx), whose prose, tables and
    figures are re-emitted section by section into the four-chapter college structure;
  * outline.txt, which says which old section goes where and holds the new writing;
  * source files quoted as code snippets.

Output: body.docx (declaration onward). finish_report.ps1 then puts the college's own starting
pages in front, updates the contents/figure/table lists in Word and saves the final report.

Formatting rules applied (from "Rules for Project Report Formatting"):
  A4; margins left 1.5 cm, top/bottom/right 1 cm; Times New Roman throughout; body 12 pt,
  justified, 1.5 line spacing; chapter headings 14 pt bold; sub-headings 12 pt bold; quotations
  and code single-spaced; page numbers centred at the bottom.
"""
import copy
import io
import re
import sys
from pathlib import Path

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
OLD = ROOT / "VoxBox_Project_Report.previous.docx"
OUT = HERE / "build" / "body.docx"
FONT = "Times New Roman"
MONO = "Courier New"


# --------------------------------------------------------------------------- styles

def force_font(style_or_run_font_element, name=FONT):
    rpr = style_or_run_font_element
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]


def setup_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    force_font(normal.element.get_or_add_rPr())
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)

    # Own style names: Word keeps the destination's definition when two files share a style name,
    # and the college's starting pages define a centred, bordered "Heading 3" of their own.
    body = styles.add_style("VB Body", 1)
    body.base_style = normal
    body.font.name = FONT
    body.font.size = Pt(12)
    force_font(body.element.get_or_add_rPr())
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.space_after = Pt(6)

    def heading(name, size, center=False, italic=False, before=12, after=6, level=0):
        st = styles.add_style(name, 1)
        st.base_style = body
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(level))
        st.element.get_or_add_pPr().append(outline)
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = italic
        st.font.color.rgb = RGBColor(0, 0, 0)
        force_font(st.element.get_or_add_rPr())
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.keep_with_next = True

    heading("VB Chapter", 14, center=True, before=0, after=18, level=0)
    heading("VB Section", 12, level=1)
    heading("VB Subsection", 12, level=2)
    heading("VB Minor", 12, italic=True, before=6, after=3, level=3)

    for name, base in (("Figure Caption", "Normal"), ("Table Caption", "Normal")):
        st = styles.add_style(name, 1)
        st.base_style = body
        st.font.size = Pt(11)
        st.font.bold = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.space_before = Pt(4)
        st.paragraph_format.space_after = Pt(10)
        st.paragraph_format.line_spacing = 1.0

    code = styles.add_style("Code Block", 1)
    code.base_style = body
    code.font.name = MONO
    code.font.size = Pt(9)
    force_font(code.element.get_or_add_rPr(), MONO)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Cm(0.3)

    for name in ("List Bullet", "List Number"):
        st = styles[name]
        st.font.name = FONT
        st.font.size = Pt(12)
        force_font(st.element.get_or_add_rPr())
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(3)


def set_margins(section):
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.0)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)
    section.footer_distance = Cm(0.5)
    section.header_distance = Cm(0.5)


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    text = OxmlElement("w:t")
    text.text = placeholder
    run._r.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def page_numbers(section, fmt, start=1):
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(footer.runs):
        r._r.getparent().remove(r._r)
    add_field(footer, "PAGE", "1")
    sect = section._sectPr
    pg = sect.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sect.append(pg)
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))


# --------------------------------------------------------------------------- old report

# Real captions put two spaces after the number; prose such as "Figure 1.1 shows…" has one.
CAPTION = re.compile(r"^(Figure|Table)\s+(\d+\.\d+)(?:\s{2,}|\t)(.*)$")
H2NUM = re.compile(r"^(\d+\.\d+)\s+(.*)$")
H3NUM = re.compile(r"^(\d+\.\d+\.\d+)\s+(.*)$")


class OldReport:
    """Splits the previous report into numbered sections of raw body elements."""

    def __init__(self, path):
        self.doc = docx.Document(str(path))
        self.sections = {}
        self.titles = {}
        current = None
        started = False
        for el in self.doc.element.body.iterchildren():
            tag = el.tag.split("}")[1]
            if tag == "p":
                p = docx.text.paragraph.Paragraph(el, self.doc)
                text = p.text.strip()
                style = self.style(p)
                if style == "Heading 1":
                    if text.startswith("1. Introduction"):
                        started = True
                    if text.startswith("Appendix") or text.startswith("18."):
                        break
                    current = None
                    continue
                if not started:
                    continue
                if style == "Heading 2":
                    m = H2NUM.match(text)
                    if m:
                        current = m.group(1)
                        self.sections[current] = []
                        self.titles[current] = m.group(2)
                        continue
            if started and current:
                self.sections[current].append(el)

    @staticmethod
    def style(p):
        try:
            return p.style.name if p.style is not None else ""
        except Exception:
            return ""

    def image_blob(self, el):
        blips = el.xpath(".//a:blip")
        if not blips:
            return None
        rid = blips[0].get(qn("r:embed"))
        return self.doc.part.related_parts[rid].blob


# --------------------------------------------------------------------------- writer

# Statements in the previous report that this release made untrue. Applied before cross-reference
# rewriting, so their own section numbers are still remapped afterwards.
STALE = [
    (re.compile(r"\b20-second\b"), "pause-aligned 16-to-28-second"),
    (re.compile(r"\btwenty-second\b"), "pause-aligned"),
    (re.compile(r"No accuracy figure is claimed anywhere in this report, because none has been measured\."),
     "Apart from the note coverage measured once in the live field test reported in the testing chapter, no accuracy figure is "
     "claimed, because no labelled corpus exists to measure one."),
    (re.compile(r"no measured accuracy of any kind"), "no measured accuracy beyond a single live coverage test"),
    (re.compile(r"and no accuracy figure is claimed anywhere\."),
     "and the only accuracy figure is the coverage measured in the live test."),
    (re.compile(r"No transcription accuracy figure\. No labelled corpus exists and no scorer has been written\."),
     "No transcription word-error-rate figure. No labelled corpus exists and no scorer has been written; note "
     "coverage was measured once, against a lecture's captions in the live field test."),
    (re.compile(r"the audio loop turns every twenty seconds regardless"),
     "the audio loop emits a chunk at each natural pause, every 16 to 28 seconds, regardless"),
    (re.compile(r"Every 20 seconds the audio pipeline emits a chunk"),
     "At the first natural pause after 16 seconds (and at most every 28 seconds) the audio pipeline emits a chunk"),
    (re.compile(r"^Every 20 s$"), "Every 16–28 s, at a pause"),
    (re.compile(r"timestamps every twenty seconds rather than"), "timestamps once per chunk (every 16 to 28 seconds) rather than"),
    (re.compile(r"\bTables \d+\.\d+ to \d+\.\d+"), "the tables of this report"),
    (re.compile(r"\b[Cc]hapter 18\b"), "the References"),
]

_SENTENCE_START = re.compile(r"(^|[.!?:]\s+)(the figure|the table|this report|the project repository|the References)\b")


class Writer:
    def __init__(self, old, emit):
        self.old = old
        self.emit = emit
        self.doc = docx.Document()
        setup_styles(self.doc)
        set_margins(self.doc.sections[0])
        self.chapter = 0
        self.h2 = 0
        self.h3 = 0
        self.fig = 0
        self.tab = 0
        self.figure_map = {}
        self.table_map = {}
        self.section_map = {}
        self.chapter_map = {}
        self.figures = []
        self.tables = []

    # numbering ---------------------------------------------------------------
    def num(self):
        return f"{self.chapter}.{self.h2}"

    LEVELS = {1: "VB Chapter", 2: "VB Section", 3: "VB Subsection", 4: "VB Minor"}

    def heading(self, text, level):
        return self.doc.add_paragraph(text, style=self.LEVELS[level])

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def chapter_heading(self, number, title):
        self.chapter = number
        self.h2 = self.h3 = self.fig = self.tab = 0
        self.heading(f"CHAPTER {number}: {title.upper()}", 1).paragraph_format.page_break_before = True

    def front_heading(self, title, new_page=True):
        p = self.heading(title, 1)
        p.paragraph_format.page_break_before = new_page

    def h2_heading(self, title):
        self.h2 += 1
        self.h3 = 0
        self.heading(f"{self.num()} {title}", 2)

    def h3_heading(self, title):
        self.h3 += 1
        self.heading(f"{self.num()}.{self.h3} {title}", 3)
        return f"{self.num()}.{self.h3}"

    # inline text -------------------------------------------------------------
    def rewrite(self, text):
        for pattern, repl in STALE:
            text = pattern.sub(repl, text)

        def fig(m):
            new = self.figure_map.get(m.group(2))
            return f"{m.group(1)} {new}" if new else "the figure"

        def tab(m):
            new = self.table_map.get(m.group(2))
            return f"{m.group(1)} {new}" if new else "the table"

        def sec(m):
            key = m.group(2)
            new = self.section_map.get(key) or self.section_map.get(".".join(key.split(".")[:2]))
            if new:
                word = m.group(1)
                # Old chapters folded into a section of a new chapter are referred to as sections.
                if "." in new and word.lower().startswith("chapter"):
                    word = "Sections" if word.endswith("s") else "Section"
                return f"{word} {new}"
            chapter = self.chapter_map.get(key.split(".")[0])
            return f"{m.group(1)} {chapter}" if chapter else "this report"

        text = re.sub(r"\b(Figures?)\s+(\d+\.\d+)\b", fig, text)
        text = re.sub(r"\b(Tables?)\s+(\d+\.\d+)\b", tab, text)
        text = re.sub(r"\b(Sections?|section|sections|Chapters?|chapter|chapters)\s+(\d+(?:\.\d+){0,2})\b", sec, text)
        text = re.sub(r"\bAppendi(?:x|ces)\s+[A-Z](?:\s*(?:and|,)\s*[A-Z])*\b", "the project repository", text)
        return _SENTENCE_START.sub(lambda m: m.group(1) + m.group(2)[0].upper() + m.group(2)[1:], text)

    def add_markup_runs(self, paragraph, text):
        # **bold**, *italic*, `code`
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                paragraph.add_run(part[2:-2]).bold = True
            elif part.startswith("`"):
                r = paragraph.add_run(part[1:-1])
                r.font.name = MONO
                r.font.size = Pt(10.5)
            elif part.startswith("*") and part.endswith("*"):
                paragraph.add_run(part[1:-1]).italic = True
            else:
                paragraph.add_run(part)

    def para(self, text, style=None, raw=False):
        p = self.doc.add_paragraph(style=style or "VB Body")
        self.add_markup_runs(p, text if raw else self.rewrite(text))
        return p

    # figures and tables ------------------------------------------------------
    def caption(self, kind, text, old_number=None, raw=False):
        if kind == "Figure":
            self.fig += 1
            number = f"{self.chapter}.{self.fig}"
            if old_number:
                self.figure_map[old_number] = number
            self.figures.append(number)
        else:
            self.tab += 1
            number = f"{self.chapter}.{self.tab}"
            if old_number:
                self.table_map[old_number] = number
            self.tables.append(number)
        p = self.doc.add_paragraph(style=f"{kind} Caption")
        p.add_run(f"{kind} {number}: {text if raw else self.rewrite(text)}")
        return number

    def picture(self, blob_or_path, width_cm=15.0, max_height_cm=17.0):
        from PIL import Image
        data = blob_or_path if isinstance(blob_or_path, bytes) else Path(blob_or_path).read_bytes()
        with Image.open(io.BytesIO(data)) as im:
            w, h = im.size
        width = min(width_cm, 17.5)
        if h / w * width > max_height_cm:
            width = max_height_cm * w / h
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(io.BytesIO(data), width=Cm(width))

    def screenshots(self, paths, caption):
        table = self.doc.add_table(rows=1, cols=len(paths))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, path in zip(table.rows[0].cells, paths):
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.add_run().add_picture(str(ROOT / path), height=Cm(10.5))
        self.caption("Figure", caption, raw=True)

    def table(self, rows, caption=None, old_number=None, widths=None, raw=False):
        if not rows:
            return
        cols = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=cols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j in range(cols):
                cell = t.cell(i, j)
                value = row[j] if j < len(row) else ""
                cp = cell.paragraphs[0]
                cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                cp.paragraph_format.line_spacing = 1.0
                cp.paragraph_format.space_after = Pt(0)
                run = cp.add_run(value if raw else self.rewrite(value))
                run.font.size = Pt(10)
                if i == 0:
                    run.bold = True
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:val"), "clear")
                    shade.set(qn("w:fill"), "D9D9D9")
                    cell._tc.get_or_add_tcPr().append(shade)
        if caption:
            self.caption("Table", caption, old_number, raw=raw)

    def code(self, lines, caption):
        for line in lines:
            p = self.doc.add_paragraph(style="Code Block")
            p.add_run(line.rstrip("\n") or " ")
        self.caption("Figure", caption, raw=True)

    # old content -------------------------------------------------------------
    def old_section(self, key, with_heading=True, title=None):
        elements = self.old.sections.get(key)
        if elements is None:
            raise KeyError(f"old section {key} not found")
        if with_heading:
            new = self.h3_heading(title or self.old.titles[key])
            self.section_map[key] = new
        else:
            self.section_map[key] = self.num()
        pending_table = None
        for el in elements:
            tag = el.tag.split("}")[1]
            if tag == "tbl":
                t = docx.table.Table(el, self.old.doc)
                rows = [[c.text.strip() for c in r.cells] for r in t.rows]
                # Merged cells repeat; collapse exact duplicates within a row.
                rows = [[v for k, v in enumerate(r) if k == 0 or v != r[k - 1]] for r in rows]
                pending_table = rows
                continue
            if tag != "p":
                continue
            p = docx.text.paragraph.Paragraph(el, self.old.doc)
            text = p.text.strip()
            style = OldReport.style(p)
            blob = self.old.image_blob(el)
            if not text and not blob:
                continue
            m = CAPTION.match(text)
            if m and m.group(1) == "Table" and pending_table is not None:
                self.table(pending_table, m.group(3), old_number=m.group(2))
                pending_table = None
                continue
            if pending_table is not None:
                self.table(pending_table)
                pending_table = None
            if m and m.group(1) == "Figure":
                self.caption("Figure", m.group(3), old_number=m.group(2))
                continue
            if blob:
                self.picture(blob)
                continue
            if style == "Heading 3":
                h = H3NUM.match(text)
                self.heading(h.group(2) if h else text, 4)
                continue
            fonts = {r.font.name for r in p.runs if r.text.strip()}
            if fonts == {"Consolas"}:
                for line in p.text.split("\n"):
                    self.doc.add_paragraph(line or " ", style="Code Block")
                continue
            is_list = bool(el.xpath("./w:pPr/w:numPr")) or style == "List Paragraph"
            out = self.doc.add_paragraph(style="List Bullet" if is_list else "VB Body")
            for r in p.runs:
                if not r.text:
                    continue
                nr = out.add_run(self.rewrite(r.text))
                nr.bold = r.bold
                nr.italic = r.italic
        if pending_table is not None:
            self.table(pending_table)


# --------------------------------------------------------------------------- outline

def extract_function(path, name, max_lines=48):
    lines = (ROOT / path).read_text(encoding="utf-8").splitlines()
    start = next(i for i, l in enumerate(lines) if re.search(rf"\b(fun|function)\s+(<[^>]+>\s*)?{re.escape(name)}\b", l))
    # include preceding KDoc/comment lines directly above
    s = start
    while s > 0 and lines[s - 1].strip().startswith(("*", "/**", "//", "@")):
        s -= 1
    out = lines[s:start]
    depth = 0
    seen = False
    for line in lines[start:]:
        out.append(line)
        depth += line.count("{") - line.count("}")
        seen = seen or "{" in line
        if seen and depth <= 0:
            break
    if len(out) > max_lines:
        out = out[:max_lines] + ["    // … (continues in the source file)"]
    indent = min((len(l) - len(l.lstrip()) for l in out if l.strip()), default=0)
    return [l[indent:] for l in out]


def run(old, emit, maps=None):
    w = Writer(old, emit)
    if maps:
        w.figure_map, w.table_map, w.section_map, w.chapter_map = maps
    lines = (HERE / "outline.txt").read_text(encoding="utf-8").splitlines()
    buffer = []
    table_rows = []

    def flush():
        text = " ".join(s.strip() for s in buffer).strip()
        buffer.clear()
        if text:
            w.para(text, raw=True)

    def flush_table(caption=None):
        if table_rows:
            w.table(list(table_rows), caption, raw=True)
            table_rows.clear()

    for raw in lines:
        line = raw.rstrip()
        if line.startswith(";"):
            continue
        if line.startswith("|"):
            flush()
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                table_rows.append(cells)
            continue
        if table_rows:
            if line.startswith("Table:"):
                flush_table(line[6:].strip())
                continue
            flush_table()
        if not line.strip():
            flush()
            continue
        if line.startswith("@front "):
            flush()
            w.front_heading(line[7:].strip())
        elif line.startswith("@frontsame "):
            flush()
            w.front_heading(line[11:].strip(), new_page=False)
        elif line.startswith("@chapter "):
            flush()
            num, title = line[9:].split(" ", 1)
            w.chapter_heading(int(num), title)
            for old_ch in re.findall(r"\[old (\d+(?:,\d+)*)\]", title):
                pass
        elif line.startswith("@chmap "):
            for pair in line[7:].split():
                a, b = pair.split("=")
                w.chapter_map[a] = f"{b}"
        elif line.startswith("## "):
            flush()
            w.h2_heading(line[3:].strip())
        elif line.startswith("### "):
            flush()
            w.h3_heading(line[4:].strip())
        elif line.startswith("#### "):
            flush()
            w.heading(line[5:].strip(), 4)
        elif line.startswith("@old "):
            flush()
            for key in line[5:].split():
                inline = key.endswith("!")
                w.old_section(key.rstrip("!"), with_heading=not inline)
        elif line.startswith("!fig "):
            flush()
            path, caption = [s.strip() for s in line[5:].split("|", 1)]
            w.picture(ROOT / path)
            w.caption("Figure", caption, raw=True)
        elif line.startswith("!shots "):
            flush()
            paths, caption = [s.strip() for s in line[7:].split("|", 1)]
            w.screenshots([p.strip() for p in paths.split(",")], caption)
        elif line.startswith("!code "):
            flush()
            spec, caption = [s.strip() for s in line[6:].split("|", 1)]
            path, name = spec.split("::")
            limit = 48
            if ":" in name:
                name, limit = name.split(":")
            w.code(extract_function(path, name, int(limit)), caption)
        elif line.startswith("- "):
            flush()
            w.para(line[2:], style="List Bullet", raw=True)
        elif re.match(r"^\d+\. ", line):
            flush()
            w.para(line, raw=True).paragraph_format.left_indent = Cm(0.6)
        elif line == "@pagebreak":
            flush()
            w.page_break()
        elif line.startswith("@field "):
            flush()
            p = w.doc.add_paragraph()
            add_field(p, line[7:], "Right-click and choose Update Field.")
        elif line == "@section-arabic":
            flush()
            new = w.doc.add_section(WD_SECTION.NEW_PAGE)
            set_margins(new)
            page_numbers(new, "decimal", 1)
        else:
            buffer.append(line)
    flush()
    flush_table()
    page_numbers(w.doc.sections[0], "lowerRoman", 1)
    return w


def main():
    old = OldReport(OLD)
    first = run(old, emit=False)
    maps = (first.figure_map, first.table_map, first.section_map, first.chapter_map)
    final = run(old, emit=True, maps=maps)
    OUT.parent.mkdir(exist_ok=True)
    final.doc.save(OUT)
    print(f"wrote {OUT} figures={len(final.figures)} tables={len(final.tables)}")


if __name__ == "__main__":
    sys.exit(main())
